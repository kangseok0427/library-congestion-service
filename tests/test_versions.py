"""Disposable synthetic Excel; exercise recovery and the actual visitor API."""
import json
import os
import subprocess
import sys
import threading
from datetime import datetime
from pathlib import Path

import pytest
import yaml
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import FileProvider, create_app
from backend.library_hours import KST
from library_etl import refresh
from library_etl.locking import data_lock
from library_etl.versions import VersionError, VersionStore
from tests.test_json_refresh import excel


def upload(store, tmp_path, value=3, **options):
    source = excel(tmp_path / 'fixture.xlsx', value=value)
    return store.upload_excel(source, source_name='합성자료.xlsx', partial_dates=[], **options)


def persisted(store):
    return (store.path.read_bytes(), store.list_versions(),
            {p.name: p.read_bytes() for p in store.directory.glob('*.json')})


@pytest.fixture
def store(tmp_path):
    return VersionStore(tmp_path / 'live.json')


def test_publish_rotate_restart_and_rollback_api(store, tmp_path):
    client = TestClient(create_app(FileProvider(store.path),
                       clock=lambda: datetime(2026, 9, 10, 12, tzinfo=KST)))
    ids, api = [], []
    for value in range(1, 6):
        result = upload(store, tmp_path, value)
        ids.append(result['version']['id'])
        assert result['validation'] == {'warning_count': 0, 'warnings': []}
        assert result['version']['source_name'] == '합성자료.xlsx'
        body = client.get('/api/v1/stats?date=2026-09-10').json()
        assert body['hourly_total_in'] == 12 * (2 * value + 1)
        api.append(body)
        assert len(store.list_versions()['versions']) == min(value, 4)
        assert len(list(store.directory.glob('*.json'))) == min(value, 4) + 1
    versions = store.list_versions()
    assert [v['id'] for v in versions['versions']] == ids[:0:-1]
    assert not (store.directory / (ids[0] + '.json')).exists()
    restarted = VersionStore(store.path)
    assert restarted.list_versions() == versions
    assert restarted.rollback(ids[1]) == {'status': 'rolled_back', 'active_version_id': ids[1]}
    assert client.get('/api/v1/stats?date=2026-09-10').json()['hourly'] == api[1]['hourly']
    assert sum(v['is_active'] for v in restarted.list_versions()['versions']) == 1
    assert VersionStore(store.path).list_versions()['active_version_id'] == ids[1]


@pytest.mark.parametrize('field,value', [('IN_08', None), ('IN_08', -1), ('IN_08', 1.5),
                                       ('IN_08', '=1+1'), ('수집일자', '2026-02-30'),
                                       ('게이트', 'other')])
def test_invalid_excel_keeps_every_committed_byte(store, tmp_path, field, value):
    upload(store, tmp_path)
    before = persisted(store)
    source = excel(tmp_path / 'bad.xlsx')
    book = load_workbook(source)
    header = [cell.value for cell in book.active[1]]
    book.active.cell(2, header.index(field) + 1).value = value
    book.save(source)
    book.close()
    with pytest.raises(VersionError) as caught:
        store.upload_excel(source, partial_dates=[])
    assert caught.value.code == 'INVALID_EXCEL' and caught.value.status_code == 422
    assert caught.value.as_dict()['error']['details'][0]['row'] == 2
    assert persisted(store) == before


@pytest.mark.parametrize('kind', ['duplicate', 'columns', 'extension', 'corrupt'])
def test_other_invalid_workbooks(store, tmp_path, kind):
    upload(store, tmp_path)
    before = persisted(store)
    source = excel(tmp_path / 'bad.xlsx', duplicate=kind == 'duplicate')
    if kind == 'columns':
        book = load_workbook(source)
        book.active.cell(1, 1).value = 'wrong'
        book.save(source)
        book.close()
    if kind == 'extension':
        source = source.rename(source.with_suffix('.xls'))
    if kind == 'corrupt':
        source.write_bytes(b'not-an-excel')
    with pytest.raises(VersionError, match='.') as caught:
        store.upload_excel(source)
    assert caught.value.code == 'INVALID_EXCEL'
    assert persisted(store) == before


@pytest.mark.parametrize('stage', ['rebuild', 'stage', 'version', 'active', 'commit'])
def test_failures_restore_all_four_versions(store, tmp_path, monkeypatch, stage):
    import library_etl.versions as module
    for value in range(4):
        upload(store, tmp_path, value)
    before = persisted(store)
    with monkeypatch.context() as patch:
        if stage == 'rebuild':
            patch.setattr(store, 'rebuild_fn', lambda *args: (_ for _ in ()).throw(OSError('rebuild')))
        if stage == 'stage':
            patch.setattr(module.os, 'fsync', lambda *args: (_ for _ in ()).throw(OSError('disk')))
        if stage in ('version', 'active'):
            original = module.os.replace
            failed = False
            def replace(source, destination):
                nonlocal failed
                destination = Path(destination)
                target = destination == store.path if stage == 'active' else destination.parent == store.directory and destination.suffix == '.json'
                if target and not failed:
                    failed = True
                    raise OSError('injected replace')
                return original(source, destination)
            patch.setattr(module.os, 'replace', replace)
        if stage == 'commit':
            patch.setattr(store, '_commit', lambda *args: (_ for _ in ()).throw(OSError('commit')))
        with pytest.raises(VersionError) as caught:
            upload(store, tmp_path, 10)
        assert caught.value.code == 'PUBLISH_FAILED'
    assert persisted(store) == before


def test_killed_worker_before_commit_recovers_on_actual_api_read(store, tmp_path):
    for value in range(4):
        upload(store, tmp_path, value)
    before = persisted(store)
    source = excel(tmp_path / 'crash.xlsx', value=20)
    code = '''
import os, sys
from library_etl.versions import VersionStore
class CrashingStore(VersionStore):
    def _commit(self, state):
        os._exit(19)
CrashingStore(sys.argv[1]).upload_excel(sys.argv[2], partial_dates=[])
'''
    result = subprocess.run([sys.executable, '-c', code, str(store.path), str(source)], timeout=30)
    assert result.returncode == 19
    assert store.path.read_bytes() != before[0]  # Killed after replacement, before metadata commit.
    client = TestClient(create_app(FileProvider(store.path),
                       clock=lambda: datetime(2026, 9, 10, 12, tzinfo=KST)))
    assert client.get('/api/v1/stats?date=2026-09-10').status_code == 200
    assert persisted(VersionStore(store.path)) == before
    upload(store, tmp_path, 21)  # Killed process did not leave a stale OS lock.


def test_cross_process_publish_rollback_and_legacy_conflict(store, tmp_path):
    version_id = upload(store, tmp_path)['version']['id']
    before = persisted(store)
    code = '''
import sys
from library_etl.locking import data_lock
with data_lock(sys.argv[1]):
    print('locked', flush=True)
    sys.stdin.readline()
'''
    worker = subprocess.Popen([sys.executable, '-c', code, str(store.path)],
                              stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True)
    try:
        assert worker.stdout.readline().strip() == 'locked'
        for operation in [lambda: upload(store, tmp_path), lambda: store.rollback(version_id),
                          lambda: store.list_versions(),
                          lambda: refresh(excel(tmp_path / 'cli.xlsx'), store.path, partial_dates=[])]:
            with pytest.raises(VersionError) as caught:
                operation()
            assert caught.value.code == 'PUBLISH_IN_PROGRESS'
            assert caught.value.status_code == 409
    finally:
        worker.communicate('\n', timeout=10)
    assert persisted(store) == before


def test_actual_concurrent_upload_and_rollback(store, tmp_path):
    from scripts.rebuild import rebuild
    version_id = upload(store, tmp_path)['version']['id']
    entered, release = threading.Event(), threading.Event()
    def paused(records, destination):
        entered.set()
        assert release.wait(10)
        return rebuild(records, destination)
    store.rebuild_fn = paused
    results = []
    worker = threading.Thread(target=lambda: results.append(upload(store, tmp_path, 10)))
    worker.start()
    try:
        assert entered.wait(10)
        with pytest.raises(VersionError) as caught:
            VersionStore(store.path).rollback(version_id)
        assert caught.value.code == 'PUBLISH_IN_PROGRESS'
        contender = excel(tmp_path / 'contender.xlsx', value=30)
        with pytest.raises(VersionError) as caught:
            VersionStore(store.path).upload_excel(contender, partial_dates=[])
        assert caught.value.code == 'PUBLISH_IN_PROGRESS'
    finally:
        release.set()
        worker.join(15)
    assert len(results) == 1
    assert store.list_versions()['active_version_id'] == results[0]['version']['id']
    store.rebuild_fn = rebuild
    store.rollback(version_id)
    assert {r['in_count'] for r in json.loads(store.path.read_bytes())} == {3, 4}


def test_history_partial_policy_and_legacy_api_are_preserved(store, tmp_path):
    upload(store, tmp_path)
    before = persisted(store)
    with pytest.raises(VersionError) as caught:
        store.upload_excel(excel(tmp_path / 'partial.xlsx'))
    assert caught.value.code == 'INVALID_EXCEL'
    assert persisted(store) == before
    source = excel(tmp_path / 'next.xlsx', day='2026-09-11')
    refresh(source, store.path, partial_dates=[])
    assert len(json.loads(store.path.read_bytes())) == 64
    client = TestClient(create_app(FileProvider(store.path), upload_token='synthetic'))
    response = client.post('/api/v1/admin/records', json=json.loads(store.path.read_bytes()),
                           headers={'Authorization': 'Bearer synthetic'})
    assert response.status_code == 409


@pytest.mark.parametrize('version_id', ['../../outside', '20260101T000000+0900', '', None])
def test_invalid_rollback_target(store, tmp_path, version_id):
    upload(store, tmp_path)
    before = persisted(store)
    with pytest.raises(VersionError) as caught:
        store.rollback(version_id)
    assert caught.value.code == 'VERSION_NOT_FOUND' and caught.value.status_code == 404
    assert persisted(store) == before


def test_original_name_never_becomes_a_storage_path(store, tmp_path):
    source = excel(tmp_path / 'safe.xlsx')
    with pytest.raises(VersionError):
        store.upload_excel(source, source_name='../../escape.xlsx')
    assert not store.path.exists()


@pytest.mark.parametrize('initial', [False, True])
def test_results_match_existing_v2_contract(store, tmp_path, initial):
    contract = yaml.safe_load(Path('contracts/openapi-v2.yaml').read_text(encoding='utf-8'))
    schemas = contract['components']['schemas']
    def validate(value, schema):
        if '$ref' in schema:
            return validate(value, schemas[schema['$ref'].split('/')[-1]])
        if 'const' in schema:
            assert value == schema['const']
        kind = schema.get('type')
        if kind == 'object':
            assert isinstance(value, dict)
            assert set(schema.get('required', [])) <= value.keys()
            if schema.get('additionalProperties') is False:
                assert value.keys() <= schema['properties'].keys()
            for key, child in value.items():
                if key in schema.get('properties', {}):
                    validate(child, schema['properties'][key])
        if kind == 'array':
            assert isinstance(value, list)
            assert len(value) <= schema.get('maxItems', len(value))
            for child in value:
                validate(child, schema['items'])
        if kind == 'string':
            import re
            assert isinstance(value, str)
            if 'pattern' in schema:
                assert re.fullmatch(schema['pattern'], value)
            if schema.get('format') == 'date-time':
                assert datetime.fromisoformat(value).utcoffset().total_seconds() == 9 * 3600
        if kind == 'integer':
            assert type(value) is int and value >= schema.get('minimum', value)
        if kind == 'boolean':
            assert type(value) is bool
    if initial:
        from datetime import date
        from scripts.rebuild import rebuild
        from scripts.sample import generate
        rebuild(generate(end=date(2026, 9, 10), days=2), store.path)
        validate(store.list_versions(), schemas['VersionList'])
    published = upload(store, tmp_path)
    for value, schema in [(published, 'UploadPublished'), (store.list_versions(), 'VersionList'),
                          (store.rollback(published['version']['id']), 'RollbackResult')]:
        validate(value, schemas[schema])
