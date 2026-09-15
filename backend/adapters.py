"""Independent v1 adapter; replace with T02 records without changing consumers."""
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .domain import DAYS, HOURS, DataError, parse_date, validate_records


def read_records(path):
    try:
        return validate_records(json.loads(Path(path).read_text(encoding='utf-8-sig')))
    except (OSError, json.JSONDecodeError) as exc:
        raise DataError('records 파일을 읽을 수 없습니다.') from exc


def from_excel(path, partial_dates=()):
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise DataError('xlsx 파일이 필요합니다.', 'EXCEL_FORMAT_ERROR')
    required = ['수집일자', '게이트', '통로ID', '전체_IN', '전체_OUT'] + [
        f'{direction}_{h:02}' for h in HOURS for direction in ('IN', 'OUT')]
    partial_dates = {parse_date(d).isoformat() for d in partial_dates}
    records, warnings, errors = [], [], []
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise DataError('Excel 파일을 열 수 없습니다.', 'EXCEL_FORMAT_ERROR') from exc
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, ())
        if any(header.count(k) != 1 for k in required):
            raise DataError('필수 컬럼 누락 또는 중복', 'EXCEL_FORMAT_ERROR')
        for row_number, values in enumerate(iterator, 2):
            row = dict(zip(header, values))
            if not any(v is not None for v in values) or row['수집일자'] == '전체':
                continue
            try:
                raw_day = row['수집일자']
                day = parse_date(raw_day.strftime('%Y-%m-%d') if isinstance(raw_day, datetime) else raw_day)
                gate = {'자료실.정문': 'front', '자료실.후문': 'back'}[row['게이트']]
                def number(key):
                    value = row[key]
                    if isinstance(value, str) and value.strip().isdigit():
                        value = int(value.strip())
                    if type(value) is not int or value < 0:
                        raise ValueError(f'{key}: 결측/음수/소수/수식/형식 오류')
                    return value
                total_in, total_out = number('전체_IN'), number('전체_OUT')
                daily = []
                for hour in HOURS:
                    daily.append(dict(date=day.isoformat(), day_of_week=DAYS[day.weekday()],
                                      gate=gate, gate_name=row['게이트'], passage_id=str(row['통로ID'] or ''),
                                      hour=hour, in_count=number(f'IN_{hour:02}'), out_count=number(f'OUT_{hour:02}'),
                                      total_in=total_in, total_out=total_out,
                                      is_partial=day.isoformat() in partial_dates, source_file=path.name))
                for field, total in [('in_count', total_in), ('out_count', total_out)]:
                    if sum(r[field] for r in daily) != total:
                        warnings.append(dict(row=row_number, code='TOTAL_MISMATCH', field=field))
                records.extend(daily)
            except (ValueError, KeyError, TypeError) as exc:
                errors.append(dict(row=row_number, message=str(exc)))
        if errors:
            raise DataError(f'Excel 검증 실패: {errors[:10]} (총 {len(errors)}건)', 'EXCEL_FORMAT_ERROR')
        return validate_records(records), dict(record_count=len(records), warnings=warnings, errors=errors)
    finally:
        workbook.close()
