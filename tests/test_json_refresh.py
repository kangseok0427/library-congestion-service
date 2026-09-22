"""Synthetic workbooks only; exercise the real parser, rebuild and HTTP API."""
import copy
import json
import os
from datetime import date, datetime
from pathlib import Path
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.app import FileProvider, create_app
from backend.domain import DataError, validate_records
from backend.library_hours import KST
from scripts.rebuild import rebuild
from scripts.sample import generate


def excel(path, value=3, bad=False, duplicate=False, day='2026-09-10'):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Data'
    header = ['수집일자', '게이트', '통로ID', '전체_IN', '전체_OUT'] + [
        f'{kind}_{hour:02}' for hour in range(8, 24) for kind in ('IN', 'OUT')]
    sheet.append(header)
    for index, gate in enumerate(('자료실.정문', '자료실.후문')):
        row = [day, gate, f'synthetic-{index}', 999, 888] + [value + index] * 32
        if bad:
            row[5] = '=1+1'
        sheet.append(row)
        if duplicate:
            sheet.append(row)
    book.save(path)
    book.close()
    return path


@pytest.fixture
def active(tmp_path):
    path = tmp_path / 'records.json'
    records = generate(end=date(2026, 9, 10), days=14)
    records[0].update(is_closed_day=None, is_low_volume=False, quality_note='retained')
    # Seed a v1 file directly: testing the new validator is a separate assertion.
    path.write_text(json.dumps(records), encoding='utf-8')
    client = TestClient(create_app(FileProvider(path), clock=lambda: datetime(2026,9,10,12,tzinfo=KST)))
    return path, client


def snapshot(path, client):
    response = client.get('/api/v1/stats?date=2026-09-10')
    assert response.status_code == 200
    forecast = client.get('/api/v1/congestion/today?date=2026-09-17')
    assert forecast.status_code == 200
    stable_forecast = forecast.json()
    stable_forecast.pop('reference_time')  # Request time is not a stored data value.
    return path.read_bytes(), path.stat().st_mtime_ns, response.json(), stable_forecast


def test_refresh_success_idempotence_history_and_api(active, tmp_path):
    from library_etl import refresh
    path, client = active
    before = snapshot(path, client)
    forecast_before = client.get('/api/v1/congestion/today?date=2026-09-17').json()['hourly']
    source = excel(tmp_path / 'new.xlsx')
    rebuild_spy = Mock(wraps=rebuild)
    report = refresh(source, path, partial_dates=[], rebuild_fn=rebuild_spy)
    after = snapshot(path, client)
    assert after[0] != before[0] and after[2]['total_in'] == 1998
    assert after[2]['hourly'][2]['out_count'] is None
    assert after[2]['hourly_total_out'] is None
    assert after[2]['total_out'] == 1776
    assert rebuild_spy.call_count == 1
    assert Path(rebuild_spy.call_args.args[1]) != path
    assert client.get('/api/v1/congestion/today?date=2026-09-17').json()['hourly'] != forecast_before
    rows = json.loads(after[0])
    assert len(rows) == len(json.loads(before[0]))
    assert any(r.get('quality_note') == 'retained' for r in rows)
    assert report['warning_counts']['HOURLY_TOTAL_MISMATCH'] == 4
    assert all(r['out_count'] is None for r in rows if r['date'] == '2026-09-10' and r['hour'] == 11)
    assert any('HOURLY_TOTAL_MISMATCH' in r.get('quality_note', '') for r in rows)
    refresh(source, path, partial_dates=[])
    repeated = json.loads(path.read_text(encoding='utf-8'))
    assert repeated == rows
    assert len({(r['date'], r['gate'], r['hour']) for r in rows}) == len(rows)
    assert not list(tmp_path.glob('.records.json.*'))


@pytest.mark.parametrize('failure', ['preprocess', 'duplicate', 'rebuild', 'replace'])
def test_failed_refresh_preserves_file_and_api(active, tmp_path, monkeypatch, failure):
    from library_etl import refresh as refresh_function
    import importlib
    module = importlib.import_module('library_etl.refresh')
    path, client = active
    before = snapshot(path, client)
    source = excel(tmp_path / 'new.xlsx', bad=failure == 'preprocess', duplicate=failure == 'duplicate')
    options = {}
    if failure == 'rebuild':
        def broken(records, destination):
            rebuild(records, destination)
            raise RuntimeError('injected rebuild failure after staging')
        options['rebuild_fn'] = broken
    if failure == 'replace':
        real_replace = os.replace
        def denied(source, destination):
            if Path(destination) == path:
                raise PermissionError('injected final replacement failure')
            return real_replace(source, destination)
        monkeypatch.setattr(module.os, 'replace', denied)
    with pytest.raises(DataError):
        refresh_function(source, path, partial_dates=[], **options)
    assert snapshot(path, client) == before
    assert not list(tmp_path.glob('.records.json.*'))


def test_v1_and_v11_quality_and_null_api(tmp_path):
    records = generate(end=date(2026, 9, 10), days=1)
    assert validate_records(records) == sorted(records, key=lambda r: (r['date'], r['hour'], r['gate']))
    for row in records:
        row.update(is_closed_day=None, is_low_volume=True, quality_note='HOURLY_TOTAL_MISMATCH')
        if row['hour'] == 11:
            row['out_count'] = None
    path = tmp_path / 'records.json'
    rebuild(records, path)
    saved = json.loads(path.read_text(encoding='utf-8'))
    assert all(r['is_low_volume'] is True and r['is_closed_day'] is None for r in saved)
    client = TestClient(create_app(FileProvider(path), clock=lambda: datetime(2026,9,10,12,tzinfo=KST)))
    data = client.get('/api/v1/stats?date=2026-09-10').json()
    assert data['hourly'][2]['out_count'] is None
    assert data['hourly_total_out'] is None
    assert data['hourly'][0]['quality_note'] == 'HOURLY_TOTAL_MISMATCH'


@pytest.mark.parametrize('field,value', [('is_closed_day', 1), ('is_low_volume', None), ('quality_note', []), ('out_count', True)])
def test_invalid_optional_types_rejected(field, value):
    records = generate(end=date(2026, 9, 10), days=1)
    records[0][field] = value
    with pytest.raises(DataError):
        validate_records(records)


def test_complete_to_partial_and_corrupt_existing_rejected(active, tmp_path):
    from library_etl import refresh
    path, client = active
    source = excel(tmp_path / 'new.xlsx')
    before = snapshot(path, client)
    with pytest.raises(DataError):
        refresh(source, path)  # Latest day conservatively partial.
    assert snapshot(path, client) == before
    path.write_text('{broken', encoding='utf-8')
    with pytest.raises(DataError):
        refresh(source, path, partial_dates=[])
    assert path.read_text(encoding='utf-8') == '{broken'


def test_provider_detects_same_size_same_timestamp_replacement(tmp_path):
    path = tmp_path / 'records.json'
    rows = generate(end=date(2026, 9, 10), days=1)
    rebuild(rows, path)
    provider = FileProvider(path)
    old = provider.get()
    stat = path.stat()
    changed = copy.deepcopy(rows)
    changed[0]['in_count'] += 1
    candidate = tmp_path / 'candidate.json'
    rebuild(changed, candidate)
    assert candidate.stat().st_size == stat.st_size
    os.utime(candidate, ns=(stat.st_atime_ns, stat.st_mtime_ns))
    os.replace(candidate, path)
    assert provider.get().records != old.records


def test_single_gate_null_propagates_without_changing_daily_total(tmp_path):
    rows = generate(end=date(2026, 9, 10), days=1)
    next(r for r in rows if r['hour'] == 11)['out_count'] = None
    service = rebuild(rows, tmp_path / 'records.json')
    data = service.stats('2026-09-10')
    assert data['hourly'][2]['out_count'] is None
    assert data['hourly_total_out'] is None
    assert data['total_out'] == sum(r['total_out'] for r in rows if r['hour'] == 8)


@pytest.mark.parametrize('field,value', [('IN_08', None), ('IN_08', -1), ('IN_08', 1.5),
                                        ('IN_08', True), ('IN_08', 'NaN'), ('OUT_12', None),
                                        ('수집일자', '2026-02-30'), ('게이트', 'other'), ('통로ID', None)])
def test_parser_reuses_strict_validation(tmp_path, field, value):
    from openpyxl import load_workbook
    from library_etl import preprocess
    source = excel(tmp_path / 'bad.xlsx')
    book = load_workbook(source)
    header = [c.value for c in book.active[1]]
    book.active.cell(2, header.index(field) + 1).value = value
    book.save(source)
    book.close()
    with pytest.raises(DataError):
        preprocess(source)


def test_sheet_and_empty_input_rejected(tmp_path):
    from library_etl import preprocess
    source = excel(tmp_path / 'book.xlsx')
    with pytest.raises(DataError):
        preprocess(source, sheet='missing')
    book = Workbook()
    book.save(source)
    book.close()
    with pytest.raises(DataError):
        preprocess(source)


def test_stage_failure_and_rebuild_mutation_preserve_live(active, tmp_path, monkeypatch):
    import importlib
    from library_etl import refresh
    module = importlib.import_module('library_etl.refresh')
    path, client = active
    before = snapshot(path, client)
    source = excel(tmp_path / 'new.xlsx')
    def mutate(records, destination):
        records[0]['in_count'] += 1
        rebuild(records, destination)
    with pytest.raises(DataError):
        refresh(source, path, partial_dates=[], rebuild_fn=mutate)
    assert snapshot(path, client) == before
    monkeypatch.setattr(module.os, 'fsync', Mock(side_effect=OSError('disk failure')))
    with pytest.raises(DataError):
        refresh(source, path, partial_dates=[])
    assert snapshot(path, client) == before
    assert not list(tmp_path.glob('.records.json.*'))


def test_competing_writer_rejected_and_first_writer_preserves_history(active, tmp_path):
    from library_etl import refresh
    path, client = active
    source = excel(tmp_path / 'first.xlsx')
    other = excel(tmp_path / 'other.xlsx', day='2026-09-11')
    def competing(records, destination):
        with pytest.raises(DataError, match='잠금'):
            refresh(other, path, partial_dates=[])
        return rebuild(records, destination)
    refresh(source, path, partial_dates=[], rebuild_fn=competing)
    refresh(other, path, partial_dates=[])
    saved = json.loads(path.read_text(encoding='utf-8'))
    assert len(saved) == 15 * 32
    assert {r['date'] for r in saved} >= {'2026-09-10', '2026-09-11'}


def test_cli_uses_library_records_and_reports_failure(tmp_path):
    import subprocess
    import sys
    source = excel(tmp_path / 'new.xlsx')
    path = tmp_path / 'live.json'
    env = {**os.environ, 'LIBRARY_RECORDS': str(path)}
    args = [sys.executable, '-m', 'library_etl', 'refresh', str(source), '--all-complete']
    success = subprocess.run(args, env=env, capture_output=True, text=True)
    assert success.returncode == 0, success.stderr
    assert json.loads(success.stdout)['integration_status'] == 'rebuilt'
    before = path.read_bytes()
    excel(source, bad=True)
    failure = subprocess.run(args, env=env, capture_output=True, text=True)
    assert failure.returncode == 2
    assert json.loads(failure.stderr)['error']['code'] == 'EXCEL_FORMAT_ERROR'
    assert path.read_bytes() == before
